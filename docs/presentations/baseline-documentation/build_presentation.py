"""Assemble the ordered source folders through the unchanged ASDLC builder."""
from __future__ import annotations

import base64
from datetime import datetime, timezone
import hashlib
import io
import json
from pathlib import Path
import re
import shutil
import sys

DECK = Path(__file__).resolve().parent
BUILD = DECK / "build"
CONSUMER = DECK.parents[2]
METHOD = CONSUMER.parent / "agentic-sdlc-governance"
SKILL = METHOD / "asdlc/skills/asdlc-presentation"
RUN = "20261014T205435Z-149579087f"
CACHE = DECK.parent / f"baseline-review-{RUN}" / "build"
sys.path[:0] = [str(CACHE / "python-deps"), str(SKILL / "src"), str(SKILL / "types")]

from PIL import Image
from pptx import Presentation
from export import build, stable_pptx


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def write_json(path, value):
    path.write_text(json.dumps(value, indent=2, ensure_ascii=False) + "\n")


def main():
    folders = sorted(p for p in DECK.iterdir() if p.is_dir() and re.fullmatch(r"\d{2}-\d{2}-.+", p.name))
    assert len(folders) == 37, "Expected all 37 ordered source folders"
    original = {}
    for folder in folders:
        assert not folder.is_symlink()
        files = sorted(folder.iterdir())
        assert len(files) == 2 and all(p.is_file() and not p.is_symlink() for p in files)
        assert (folder / "speaker-notes.md").is_file()
        assert len(list(folder.glob("visual.*"))) == 1
        original.update({p: digest(p) for p in files})

    cache_manifest = CACHE / "mermaid-render-manifest.json"
    renderer = json.loads(cache_manifest.read_text())
    cached = {d["id"]: d for d in renderer["diagrams"]}
    assert len(cached) == 24
    assert not BUILD.is_symlink()
    BUILD.mkdir(exist_ok=True)
    stage = BUILD / "staging" / DECK.name
    stage.mkdir(parents=True, exist_ok=True)
    rendered = BUILD / "rendered"
    rendered.mkdir(exist_ok=True)
    records = []

    for number, folder in enumerate(folders, 1):
        visual = next(folder.glob("visual.*"))
        notes = folder / "speaker-notes.md"
        target = stage / f"{number:02d}-{folder.name}"
        target.mkdir(exist_ok=True)
        shutil.copyfile(notes, target / notes.name)
        record = {"slide": number, "folder": folder.name,
                  "visual_source": str(visual.relative_to(CONSUMER)),
                  "notes_source": str(notes.relative_to(CONSUMER)),
                  "visual_sha256": original[visual], "notes_sha256": original[notes],
                  "raw_notes": notes.read_text(), "cached_render": None}
        if visual.suffix == ".svg":
            shutil.copyfile(visual, target / "visual.svg")
        else:
            assert visual.suffix == ".mmd"
            job = cached[folder.name]
            assert job["source_sha256"] == original[visual], f"Stale render: {folder.name}"
            png, svg = Path(job["output"]), Path(job["svg"])
            assert digest(png) == job["output_sha256"] and digest(svg) == job["svg_sha256"]
            local_png = rendered / f"{folder.name}.png"
            shutil.copyfile(png, local_png)
            shutil.copyfile(svg, rendered / f"{folder.name}.svg")
            with Image.open(local_png) as image:
                width, height = image.size
            payload = base64.b64encode(local_png.read_bytes()).decode("ascii")
            wrapper = (f'<svg xmlns="http://www.w3.org/2000/svg" '
                       f'xmlns:xlink="http://www.w3.org/1999/xlink" width="{width}" height="{height}" '
                       f'viewBox="0 0 {width} {height}"><image width="{width}" height="{height}" '
                       f'xlink:href="data:image/png;base64,{payload}"/></svg>\n')
            (target / "visual.svg").write_text(wrapper)
            record["cached_render"] = str(local_png.relative_to(BUILD))
            record["cached_render_sha256"] = job["output_sha256"]
        records.append(record)

    expected_staging = {f"{i:02d}-{p.name}" for i, p in enumerate(folders, 1)}
    assert {p.name for p in stage.iterdir() if p.is_dir() and p.name != "build"} == expected_staging
    result = build(stage, "intro")
    assert result.slide_count == 37
    presentation = Presentation(result.pptx)
    for slide, record in zip(presentation.slides, records):
        picture = slide.shapes[0]
        if record["cached_render"]:
            # Keep the full-resolution verified render for PowerPoint zooming.
            image_path = BUILD / record["cached_render"]
            with Image.open(image_path) as image:
                width, height = image.size
            scale = min(presentation.slide_width / width, presentation.slide_height / height)
            w, h = round(width * scale), round(height * scale)
            old = picture
            picture = slide.shapes.add_picture(str(image_path),
                (presentation.slide_width-w)//2, (presentation.slide_height-h)//2, width=w, height=h)
            old_id = old._pic.blipFill.blip.rEmbed
            old._element.getparent().remove(old._element)
            if old_id != picture._pic.blipFill.blip.rEmbed:
                slide.part.drop_rel(old_id)
        picture.name = record["folder"]
        picture._element.nvPicPr.cNvPr.set("descr", record["visual_source"])
        slide.notes_slide.notes_text_frame.text = record["raw_notes"]
        record["embedded_image_sha256"] = hashlib.sha256(picture.image.blob).hexdigest()

    presentation.core_properties.title = "Speaker Coach baseline documentation"
    presentation.core_properties.subject = f"{RUN}; CAPTURED_NEEDS_HUMAN_REVIEW"
    final = BUILD / "baseline-documentation.pptx"
    final.write_bytes(stable_pptx(presentation))
    reread = Presentation(final)
    assert len(reread.slides) == len(records) == 37
    for slide, record in zip(reread.slides, records):
        assert len(slide.shapes) == 1
        picture = slide.shapes[0]
        assert picture.name == record["folder"]
        assert slide.notes_slide.notes_text_frame.text == record["raw_notes"]
        assert hashlib.sha256(picture.image.blob).hexdigest() == record["embedded_image_sha256"]
        assert picture.left >= 0 and picture.top >= 0
        assert picture.left + picture.width <= reread.slide_width
        assert picture.top + picture.height <= reread.slide_height
        width, height = picture.image.size
        assert abs((picture.width / picture.height) / (width / height) - 1) < 0.001
        assert not any((picture.crop_left, picture.crop_right, picture.crop_top, picture.crop_bottom))
        record["notes_round_trip_exact"] = True
    assert all(digest(p) == value for p, value in original.items()), "Source changed during build"

    (BUILD / "complete-speaker-notes.md").write_text("\n\n".join(
        f'# Slide {r["slide"]:02d} — {r["folder"]}\n\n{r["raw_notes"]}' for r in records) + "\n")
    for record in records:
        del record["raw_notes"]
    write_json(BUILD / "coverage-manifest.json", {
        "presentation": DECK.name, "source_run": RUN, "status": "CAPTURED_NEEDS_HUMAN_REVIEW",
        "built_at": datetime.now(timezone.utc).isoformat(), "slide_count": 37,
        "ordering": "Lexicographic order of complete stage-topic source folder names",
        "narration_profile": "intro; complete original notes restored verbatim",
        "pptx_sha256": digest(final), "source_files_unchanged": True,
        "notes_round_trip_exact": True, "image_geometry_verified": True,
        "render_cache_manifest": str(cache_manifest), "render_cache_manifest_sha256": digest(cache_manifest),
        "rendering_compatibility_encodings": renderer["rendering_compatibility_encodings"],
        "limitations": ["PowerPoint visuals are images; original SVG/Mermaid sources remain editable.",
                        "Dense complete diagrams require zoom for detailed reading."], "slides": records})
    shutil.copyfile(result.manifest, BUILD / "shared-build-manifest.json")
    shutil.copyfile(stage / "build/preview.png", BUILD / "preview.png")
    print(json.dumps({"pptx": str(final), "slides": 37, "exact_notes": True,
                      "source_files_unchanged": True, "geometry_verified": True}))


if __name__ == "__main__":
    main()
