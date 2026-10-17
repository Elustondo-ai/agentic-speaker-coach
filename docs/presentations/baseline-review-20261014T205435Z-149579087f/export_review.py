"""Deck-local faithful export; original baseline and shared framework are read-only."""
from __future__ import annotations
import base64
from datetime import datetime, timezone
import hashlib
import io
import json
import math
from pathlib import Path
import shutil
import sys
from xml.sax.saxutils import escape

DECK = Path(__file__).resolve().parent
BUILD = DECK / "build"
CONSUMER = DECK.parents[2]
METHOD = CONSUMER.parent / "agentic-sdlc-governance"
SKILL = METHOD / "asdlc/skills/asdlc-presentation"
RUN = "20261014T205435Z-149579087f"
RECEIPT = CONSUMER / f"docs/asdlc/00-governance/logs/{RUN}__baseline-run.json"
sys.path[:0] = [str(BUILD/"python-deps"), str(SKILL/"src"), str(SKILL/"types")]
import ctypes
ctypes.CDLL("/opt/homebrew/lib/libcairo.2.dylib", mode=ctypes.RTLD_GLOBAL)
from PIL import Image, ImageDraw, ImageFont, ImageOps
from pptx import Presentation
from export import build, stable_pptx

STAGES = {"00":"Governance", "01":"Context and Domain Analysis", "02":"Intent",
          "03":"System Model and Use Cases", "04":"Benchmarks, Test Strategy and Success Criteria",
          "05":"High-Level Solution Design and Architecture", "08":"Low-Level Design",
          "11":"Operational Reality, DevOps and User Acceptance"}
FONT_PATH = "/System/Library/Fonts/Supplemental/Arial.ttf"
FONT = ImageFont.truetype(FONT_PATH, 32)
STAMP = datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def write_json(path, value):
    path.write_text(json.dumps(value, indent=2, ensure_ascii=False)+"\n")


def load_topics():
    receipt = json.loads(RECEIPT.read_text())
    assert receipt["run_id"] == RUN and receipt["status"] == "CAPTURED_NEEDS_HUMAN_REVIEW"
    artifacts = {e["path"]:e for invocation in receipt["invocations"] for e in invocation["artifacts"]}
    assert len(artifacts) == 48
    assert {str(CONSUMER/p) for p in artifacts} == set(receipt["written_artifacts"])
    for rel, entry in artifacts.items():
        path = CONSUMER/rel
        assert RUN in path.parts and "baselines" in path.parts
        assert not path.is_symlink() and path.is_file() and path.stat().st_nlink == 1
        assert digest(path) == entry["sha256"], rel
    topics = []
    for folder in sorted({Path(p).parent for p in artifacts}):
        stage = folder.parts[2][:2]
        notes, visual = CONSUMER/folder/"speaker-notes.md", CONSUMER/folder/"visual.mmd"
        assert str(notes.relative_to(CONSUMER)) in artifacts and str(visual.relative_to(CONSUMER)) in artifacts
        raw = notes.read_text()
        title = next(line[2:] for line in raw.splitlines() if line.startswith("# "))
        topics.append(dict(id=f"{stage}/{folder.name}", stage=stage, topic=folder.name, title=title,
                           notes=notes, visual=visual, raw=raw, notes_sha256=digest(notes),
                           visual_sha256=digest(visual), key=f"{stage}-{folder.name}"))
    assert len(topics)==24 and all(sum(t["stage"]==s for t in topics)==3 for s in STAGES)
    return receipt, topics


def text(x, y, value, size=32, color="#172A3A", weight="normal"):
    return (f'<text x="{x}" y="{y}" font-family="Arial" font-size="{size}" '
            f'font-weight="{weight}" fill="{color}">{escape(value)}</text>')


def wrap(value, font=FONT, width=1760):
    """Return exact source spans; only visual line wrapping is added."""
    pieces, offset = [], 0
    for physical in value.splitlines(keepends=True):
        end = len(physical.rstrip("\r\n"))
        if not end:
            pieces.append((offset,offset+len(physical),""))
        else:
            start = 0
            while start < end:
                stop = start+1
                while stop<=end and font.getlength(physical[start:stop])<=width:
                    stop+=1
                stop-=1
                if stop<end:
                    space=physical.rfind(" ",start,stop+1)
                    if space>start: stop=space+1
                assert stop>start
                last=len(physical) if stop==end else stop
                pieces.append((offset+start,offset+last,physical[start:stop].rstrip()))
                start=stop
        offset+=len(physical)
    assert "".join(value[a:b] for a,b,_ in pieces)==value
    return pieces


def frame(stage, title, kind, number):
    stage_name=f"{stage} — {STAGES[stage]}" if stage else "Speaker Coach | Captured baseline review"
    svg=['<svg xmlns="http://www.w3.org/2000/svg" xmlns:xlink="http://www.w3.org/1999/xlink" width="1920" height="1080" viewBox="0 0 1920 1080">',
         '<rect width="1920" height="1080" fill="white"/>',
         '<rect width="1920" height="74" fill="#102E43"/>',
         text(64,49,stage_name,30,"white"),text(64,122,kind,23,"#276779","bold")]
    lines=wrap(title,ImageFont.truetype(FONT_PATH,44),1780)
    assert len(lines)<=2
    svg += [text(64,176+48*i,line,44,weight="bold") for i,(_,_,line) in enumerate(lines)]
    svg += ['<line x1="64" y1="1016" x2="1856" y2="1016" stroke="#D6E0E5"/>',
            text(64,1056,f"{RUN}  |  CAPTURED_NEEDS_HUMAN_REVIEW",20,"#526674"),
            text(1810,1056,str(number),23,"#526674")]
    return svg


def png_image(data,x,y,w,h):
    encoded=base64.b64encode(data).decode()
    return (f'<image x="{x}" y="{y}" width="{w}" height="{h}" preserveAspectRatio="xMidYMid meet" '
            f'xlink:href="data:image/png;base64,{encoded}"/>')


def main():
    receipt, topics=load_topics()
    if "--prepare" in sys.argv:
        (BUILD/"rendered").mkdir(parents=True,exist_ok=True)
        write_json(BUILD/"mermaid-jobs.json",[
            dict(id=t["key"],source=str(t["visual"]),
                 output=str(BUILD/"rendered"/(t["key"]+".png")),
                 svg=str(BUILD/"rendered"/(t["key"]+".svg"))) for t in topics])
        print("Prepared 24 verified Mermaid jobs from the immutable receipt.")
        return
    originals={str(RECEIPT):digest(RECEIPT)}
    for t in topics:
        originals.update({str(t["notes"]):t["notes_sha256"],str(t["visual"]):t["visual_sha256"]})
    render_report=json.loads((BUILD/"mermaid-render-manifest.json").read_text())
    renders={r["id"]:r for r in render_report["diagrams"]}
    assert len(renders)==24, "Every original diagram must be rendered; no silent omission."
    staged=BUILD/"staging"/DECK.name
    if staged.exists(): shutil.rmtree(staged)  # this adapter's derived files only
    staged.mkdir(parents=True)
    slides,coverage,geometry=[],[],[]

    def add(stage,title,kind,body,notes,record):
        number=len(slides)+1
        folder=staged/f"{number:03d}-{record['kind']}"
        folder.mkdir()
        (folder/"visual.svg").write_text("\n".join(frame(stage,title,kind,number)+body+["</svg>"]))
        (folder/"speaker-notes.md").write_text(notes)
        slides.append(dict(number=number,folder=str(folder),notes=notes,title=title,**record))
        return number

    scope=[
        "Source project: agentic-speaker-coach",
        f"Exact baseline run: {RUN}",
        "Final receipt: CAPTURED_NEEDS_HUMAN_REVIEW",
        "24 topics / 48 source artifacts • 8 stage groups • faithful review export",
        "Governance first, then stages 01, 02, 03, 04, 05, 08 and 11.",
        "Each original visual is followed by complete captured notes for on-screen review.",
        "Magnified detail views are labelled supplements to the complete visual.",
        "Captured observations, claims, uncertainties and contradictions are unchanged.",
        "Governance topic 02 records a pre-synthesis snapshot, preserved as captured.",
        "Final receipt facts here are added navigation, not a rewrite of that snapshot.",
        "This deck does not grant approval or establish lifecycle completion."]
    add(None,"Speaker Coach — complete baseline for Pablo's review","ADDED NAVIGATION — EXPORT SCOPE",
        [text(70,290+i*58,line,30) for i,line in enumerate(scope)],
        f"AI-MODIFIED - {STAMP}\n# Added scope navigation\n"+"\n".join(scope)+"\n",
        dict(kind="navigation",source_kind="added_navigation"))
    for stage in STAGES:
        group=[t for t in topics if t["stage"]==stage]
        body=[text(70,300,"Captured topics in original within-stage order",34,weight="bold")]
        for i,t in enumerate(group):
            body += [text(70,400+i*155,t["topic"],32,"#276779","bold"),text(70,450+i*155,t["title"],32)]
        add(stage,"Stage topic index","ADDED NAVIGATION — NO NEW BASELINE CLAIMS",body,
            f"AI-MODIFIED - {STAMP}\n# Added stage navigation\n"+"\n".join(t["id"]+" — "+t["title"] for t in group)+"\n",
            dict(kind="navigation",source_kind="added_navigation",stage=stage))
        for topic in group:
            rendered=renders[topic["key"]]
            assert rendered["source_sha256"]==topic["visual_sha256"]
            data=Path(rendered["output"]).read_bytes()
            image=Image.open(io.BytesIO(data)).convert("RGB")
            iw,ih=image.size
            aw,ah=1780,750
            scale=min(aw/iw,ah/ih);w,h=iw*scale,ih*scale
            number=add(stage,topic["title"],f"CAPTURED VISUAL — {topic['id']} — COMPLETE DIAGRAM",
                       [png_image(data,(1920-w)/2,240+(ah-h)/2,w,h)],topic["raw"],
                       dict(kind="visual",source_kind="captured_visual",topic_id=topic["id"],image_ratio=iw/ih))
            entry=dict(topic_id=topic["id"],stage=stage,stage_name=STAGES[stage],original_topic_id=topic["topic"],
                       title=topic["title"],notes_source=str(topic["notes"].relative_to(CONSUMER)),
                       visual_source=str(topic["visual"].relative_to(CONSUMER)),notes_sha256=topic["notes_sha256"],
                       visual_sha256=topic["visual_sha256"],rendered_visual_sha256=rendered["output_sha256"],
                       rendered_visual=str(Path(rendered["output"]).relative_to(BUILD)),
                       visual_slide=number,diagram_detail_slides=[],text_slides=[])
            # Keep the complete visual. Labelled overlapping detail views make dense
            # labels readable; no original diagram is cropped or replaced.
            minfont=min(m["approximate_rendered_font_size_px"] for m in rendered["text_metrics"])
            magnify=max(scale,28/minfont)
            cols=max(1,math.ceil(iw*magnify/aw)); rows=max(1,math.ceil(ih*magnify/ah))
            entry["overview_min_font_px"]=minfont*scale
            entry["detail_min_font_px"]=minfont*magnify
            if cols*rows>1:
                tw,th=min(iw,aw/magnify),min(ih,ah/magnify)
                boxes=rendered["text_metrics"]
                def contains(view,box):
                    x,y,vw,vh=view
                    return (box["x"]>=x-2 and box["y"]>=y-2
                            and box["x"]+box["width"]<=x+vw+2
                            and box["y"]+box["height"]<=y+vh+2)
                viewports=[]
                for row in range(rows):
                    for col in range(cols):
                        x=0 if cols==1 else (iw-tw)*col/(cols-1)
                        y=0 if rows==1 else (ih-th)*row/(rows-1)
                        candidate=[x,y,tw,th]
                        if any(contains(candidate,box) for box in boxes):
                            viewports.append(candidate)
                # Omit blank/fragments-only grid tiles. Add a label-centered
                # viewport where a grid boundary split a complete label.
                for box in boxes:
                    if not any(contains(view,box) for view in viewports):
                        vw,vh=max(tw,box["width"]+32),max(th,box["height"]+32)
                        x=max(0,min(iw-vw,box["x"]+box["width"]/2-vw/2))
                        y=max(0,min(ih-vh,box["y"]+box["height"]/2-vh/2))
                        viewports.append([x,y,min(iw,vw),min(ih,vh)])
                covered=[{i for i,box in enumerate(boxes) if contains(view,box)} for view in viewports]
                # A sparse view showing only a cluster heading adds nothing
                # when a neighbouring view includes all its labels and more.
                viewports=[view for i,view in enumerate(viewports)
                           if not any(covered[i]<other or (covered[i]==other and j<i)
                                      for j,other in enumerate(covered) if j!=i)]
                assert all(any(contains(view,box) for view in viewports) for box in boxes)
                entry["all_diagram_labels_have_whole_detail_view"]=True
                for index,(x,y,vw,vh) in enumerate(viewports,1):
                    crop=image.crop((round(x),round(y),round(x+vw),round(y+vh)))
                    output=io.BytesIO();crop.save(output,"PNG")
                    detail_scale=min(aw/crop.width,ah/crop.height,magnify)
                    cw,ch=crop.width*detail_scale,crop.height*detail_scale
                    n=add(stage,topic["title"],
                          f"ADDED MAGNIFICATION — {topic['id']} — VIEW {index} OF {len(viewports)}; FULL VISUAL: SLIDE {number}",
                          [png_image(output.getvalue(),(1920-cw)/2,240+(ah-ch)/2,cw,ch)],topic["raw"],
                          dict(kind="diagram-detail",source_kind="captured_visual_magnification",topic_id=topic["id"],
                               complete_visual_slide=number,pixel_viewport=[x,y,vw,vh]))
                    entry["diagram_detail_slides"].append(n)
            pieces=wrap(topic["raw"])
            page_count=math.ceil(len(pieces)/18)
            # Balanced pages avoid exporting a lone trailing run/metadata line.
            pages=[pieces[round(i*len(pieces)/page_count):round((i+1)*len(pieces)/page_count)]
                   for i in range(page_count)]
            for i,page in enumerate(pages):
                body=[]
                for j,(_,_,line) in enumerate(page):
                    body.append(text(70,265+j*40,line,32))
                    geometry.append(dict(topic=topic["id"],width=FONT.getlength(line),max_width=1780,baseline=265+j*40))
                n=add(stage,topic["title"],f"CAPTURED NOTES — {topic['id']} — PART {i+1} OF {len(pages)} (VERBATIM; WRAPPED FOR DISPLAY)",
                      body,topic["raw"],dict(kind="notes",source_kind="captured_notes",topic_id=topic["id"],
                      notes_span=[page[0][0],page[-1][1]],display_lines=[p[2] for p in page]))
                entry["text_slides"].append(dict(slide=n,start_character=page[0][0],end_character=page[-1][1]))
            assert "".join(topic["raw"][p["start_character"]:p["end_character"]] for p in entry["text_slides"])==topic["raw"]
            entry["all_slides"]=[number]+entry["diagram_detail_slides"]+[p["slide"] for p in entry["text_slides"]]
            coverage.append(entry)
    result=build(staged,profile="intro")
    presentation=Presentation(result.pptx)
    assert len(presentation.slides)==len(slides)
    for slide,info in zip(presentation.slides,slides):
        slide.notes_slide.notes_text_frame.text=info["notes"]
    presentation.core_properties.title=f"Speaker Coach baseline review — {RUN}"
    presentation.core_properties.subject="Faithful captured baseline; CAPTURED_NEEDS_HUMAN_REVIEW"
    final=BUILD/f"{DECK.name}.pptx";final.write_bytes(stable_pptx(presentation))
    reread=Presentation(final)
    assert all(s.notes_slide.notes_text_frame.text==info["notes"] for s,info in zip(reread.slides,slides))
    preview_dir=BUILD/"slides"
    if preview_dir.exists(): shutil.rmtree(preview_dir)
    preview_dir.mkdir()
    for obsolete in BUILD.glob("contact-*.png"): obsolete.unlink()
    thumbnails=[]
    for index,slide in enumerate(reread.slides,1):
        picture=slide.shapes[0];png=picture.image.blob
        (preview_dir/f"slide-{index:03d}.png").write_bytes(png)
        im=Image.open(io.BytesIO(png)).convert("RGB")
        assert abs(im.width/im.height-16/9)<0.003
        assert abs(picture.width/picture.height-im.width/im.height)<0.003
        thumbnails.append(ImageOps.pad(im,(480,270),color="white"))
    for start,count,target in [(0,len(slides),"preview.png")]+[(s,min(12,len(slides)-s),f"contact-{s+1:03d}-{min(s+12,len(slides)):03d}.png") for s in range(0,len(slides),12)]:
        sheet=Image.new("RGB",(1500,305*math.ceil(count/3)),"#E6EAEE");draw=ImageDraw.Draw(sheet)
        for j in range(count):
            x,y=(j%3)*500+10,(j//3)*305+10
            sheet.paste(thumbnails[start+j],(x,y));draw.text((x,y+274),f"{start+j+1:03d}",fill="#172A3A")
        sheet.save(BUILD/target)
    compiled=[f"AI-MODIFIED - {STAMP}","# Complete captured baseline notes","",f"{RUN}; CAPTURED_NEEDS_HUMAN_REVIEW.",""]
    for t,e in zip(topics,coverage):
        compiled += [f"## {t['id']} — {STAGES[t['stage']]}","Slides: "+", ".join(map(str,e["all_slides"])),"",t["raw"],""]
    (BUILD/"complete-captured-notes.md").write_text("\n".join(compiled))
    html=["<!doctype html><html><head><meta charset='utf-8'><title>Speaker Coach baseline review</title>",
          "<style>body{font:18px Arial;background:#e6eaee;color:#172a3a;margin:24px}figure{max-width:1440px;margin:30px auto}img{width:100%;height:auto}figcaption{padding:10px}</style></head><body>",
          "<h1>Speaker Coach — captured baseline for human review</h1>",
          f"<p>{RUN} · {len(slides)} slides · 24 topics · CAPTURED_NEEDS_HUMAN_REVIEW</p>"]
    for s in slides:
        n=s["number"];html.append(f"<figure id='slide-{n}'><img loading='lazy' src='slides/slide-{n:03d}.png' alt='{escape(s['title'])}'><figcaption>{n:03d} · {escape(s['source_kind'])} · {escape(s['title'])}</figcaption></figure>")
    (BUILD/"preview.html").write_text("\n".join(html)+"</body></html>")
    assert all(digest(Path(p))==expected for p,expected in originals.items())
    report=dict(run_id=RUN,status=receipt["status"],source_project="agentic-speaker-coach",
                receipt=str(RECEIPT.relative_to(CONSUMER)),receipt_sha256=originals[str(RECEIPT)],
                source_topic_count=24,source_artifact_count=48,slide_count=len(slides),
                ordering="Lifecycle grouping: 00,01,02,03,04,05,08,11; original topic01/02/03 within each.",
                navigation_slide_count=sum(s["kind"]=="navigation" for s in slides),diagram_overview_count=24,
                diagram_detail_count=sum(s["kind"]=="diagram-detail" for s in slides),
                text_slide_count=sum(s["kind"]=="notes" for s in slides),topics=coverage,
                checks=dict(receipt_hashes_match=True,all_topics_represented=True,notes_exact_round_trip=True,
                            full_notes_on_every_captured_slide=True,complete_on_screen_text_spans=True,
                            original_visuals_preserved_on_overviews=True,aspect_ratios_preserved=True,
                            original_baseline_and_receipt_unchanged=True),
                notes_policy="Exact original text, including metadata/title/references/markers, restored after shared build.",
                mermaid_render_manifest="mermaid-render-manifest.json",
                rendering_compatibility_encodings=render_report["rendering_compatibility_encodings"],
                limitations=["Raster-backed visuals, not editable PowerPoint diagram shapes.",
                             "Detail slides are labelled overlapping magnifications; full original overviews remain.",
                             "Historical source/context references preserved verbatim, not rewritten as live links.",
                             "Export does not validate claims or grant lifecycle completion or approval."],
                pptx=final.name,pptx_sha256=digest(final),created_at=STAMP)
    write_json(BUILD/"coverage-manifest.json",report)
    write_json(BUILD/"slide-map.json",[{k:v for k,v in s.items() if k!="notes"} for s in slides])
    write_json(BUILD/"layout-checks.json",dict(text_geometry=geometry,all_within_bounds=all(g["width"]<=g["max_width"] and g["baseline"]<1000 for g in geometry)))
    shutil.copy2(result.manifest,BUILD/"shared-build-manifest.json")
    print(json.dumps({k:report[k] for k in ("slide_count","source_topic_count","source_artifact_count","diagram_detail_count","text_slide_count","pptx")},indent=2))


if __name__=="__main__":
    main()
